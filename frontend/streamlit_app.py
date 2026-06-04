import streamlit as st
import requests
import json
import zipfile
import re
import os
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from requests.exceptions import RequestException
from fpdf import FPDF
from docx import Document
from io import BytesIO

API_URL = os.getenv("API_URL", "http://backend:8000")

st.set_page_config(page_title="AI Resume Optimizer", layout="wide")
st.title("🚀 AI Resume Optimizer")


# ===========================
# INPUTS
# ===========================
resume_file = st.file_uploader("Upload Resume", type=["pdf", "docx"])
cover_file = st.file_uploader("Upload Cover Letter", type=["pdf", "docx"])
job_requirements = st.text_area("Job Requirements")
github_url = st.text_input("GitHub Profile URL")


# ===========================
# SESSION STATE
# ===========================
if "result" not in st.session_state:
    st.session_state.result = None


# ===========================
# JSON HELPERS
# ===========================
def extract_json(text):
    if not isinstance(text, str):
        return {}

    text = text.replace("```json", "").replace("```", "").strip()

    try:
        return json.loads(text)
    except Exception:
        pass

    try:
        match = re.search(r"\{.*\}", text, re.DOTALL)
        if match:
            return json.loads(match.group())
    except Exception:
        pass

    return {}


def get_ats(result):
    ats_block = result.get("ats_optimization", {})

    if not ats_block:
        return {}

    if isinstance(ats_block, dict) and "ats_resume" in ats_block:
        return ats_block

    raw = ats_block.get("raw_output", "") if isinstance(ats_block, dict) else ""

    if isinstance(raw, str):
        return extract_json(raw)

    return {}


def clean_text(text):
    if not text:
        return ""
    return str(text).replace("…", "...").replace("—", "-").replace("–", "-")


# ===========================
# DISPLAY HELPERS
# ===========================
def bullet_list(items):
    if not isinstance(items, list):
        return ""
    return "\n".join([f"- {item}" for item in items if item])


def contact_line(resume):
    parts = []

    for key in ["location", "email", "phone", "linkedin", "github"]:
        value = resume.get(key)
        if value:
            parts.append(value)

    return " | ".join(parts)


def section_header(title):
    st.markdown(
        f"""
        <h3 style="margin-top:0.7rem; margin-bottom:0.2rem;">
            {title}
        </h3>
        """,
        unsafe_allow_html=True
    )


def centered_header(resume):
    name = resume.get("name", "")
    title = resume.get("title", "")
    contact = contact_line(resume)

    if name:
        st.markdown(
            f"<h2 style='text-align:center; margin-bottom:0; margin-top:0;'>{name}</h2>",
            unsafe_allow_html=True
        )

    if title:
        st.markdown(
            f"<p style='text-align:center; font-weight:600; margin:0;'>{title}</p>",
            unsafe_allow_html=True
        )

    if contact:
        st.markdown(
            f"<p style='text-align:center; margin:0;'>{contact}</p>",
            unsafe_allow_html=True
        )

    st.markdown("---")


def render_skills(skills):
    if isinstance(skills, list):
        for skill_group in skills:
            if isinstance(skill_group, dict):
                category = skill_group.get("category", "")
                items = skill_group.get("items", [])

                if category and items:
                    st.markdown(f"• **{category}:** {', '.join(items)}")
                elif category:
                    st.markdown(f"• **{category}**")
            elif isinstance(skill_group, str):
                st.markdown(f"- {skill_group}")

    elif isinstance(skills, dict):
        for category, items in skills.items():
            label = category.replace("_", " ").title()
            if isinstance(items, list):
                st.markdown(f"• **{label}:** {', '.join(items)}")
            else:
                st.markdown(f"• **{label}:** {items}")


def clean_achievements(description, achievements):
    if not isinstance(achievements, list):
        return []

    desc = str(description or "").strip().lower()
    unique = []
    seen = set()

    for item in achievements:
        item_clean = str(item).strip()
        item_key = item_clean.lower()

        if not item_clean:
            continue

        if desc and item_key == desc:
            continue

        if item_key in seen:
            continue

        seen.add(item_key)
        unique.append(item_clean)

    return unique


def render_resume(resume):
    centered_header(resume)

    section_header("Summary")
    st.markdown(resume.get("summary", ""))

    st.markdown("---")

    section_header("Technologies & Skills")
    render_skills(resume.get("skills", []))


    st.markdown("---")

    section_header("Projects")
    for project in resume.get("projects", []):
        if isinstance(project, dict):
            name = project.get("name", "")
            description = project.get("description", "")
            achievements = clean_achievements(description, project.get("achievements", []))

            if name:
                st.markdown(f"**{name}**")

            if description and not achievements:
                st.markdown(description)

            if achievements:
                st.markdown(bullet_list(achievements))

            st.markdown("")

    st.markdown("---")

    section_header("Education")
    for edu in resume.get("education", []):
        if isinstance(edu, dict):
            text = " | ".join([
                x for x in [
                    edu.get("degree", ""),
                    edu.get("university", ""),
                    edu.get("year", "")
                ] if x
            ])
            st.markdown(f"- {text}")

    st.markdown("---")

    section_header("Certifications")
    for cert in resume.get("certifications", []):
        if isinstance(cert, dict):
            name = cert.get("name", "")
            desc = cert.get("description", "")

            if name and desc:
                st.markdown(f"- **{name}**: {desc}")
            elif name:
                st.markdown(f"- **{name}**")
            elif desc:
                st.markdown(f"- {desc}")


# ===========================
# PDF GENERATOR
# ===========================
def resume_to_plain_text(resume):
    lines = []

    if resume.get("name"):
        lines.append(resume.get("name"))
    if resume.get("title"):
        lines.append(resume.get("title"))

    contact = contact_line(resume)
    if contact:
        lines.append(contact)

    lines.append("")
    lines.append("SUMMARY")
    lines.append(resume.get("summary", ""))

    lines.append("")
    lines.append("TECHNOLOGIES & SKILLS")

    skills = resume.get("skills", [])

    if isinstance(skills, list):
        for skill_group in skills:
            if isinstance(skill_group, dict):
                category = skill_group.get("category", "")
                items = skill_group.get("items", [])
                if category and items:
                    lines.append(f"- {category}: {', '.join(items)}")
            else:
                lines.append(f"- {skill_group}")

    elif isinstance(skills, dict):
        for category, items in skills.items():
            label = category.replace("_", " ").title()
            if isinstance(items, list):
                lines.append(f"- {label}: {', '.join(items)}")
            else:
                lines.append(f"- {label}: {items}")

    lines.append("")
    lines.append("PROJECTS")
    for project in resume.get("projects", []):
        if isinstance(project, dict):
            name = project.get("name", "")
            description = project.get("description", "")
            achievements = clean_achievements(description, project.get("achievements", []))

            if name:
                lines.append(name)

            if description and not achievements:
                lines.append(description)

            for achievement in achievements:
                lines.append(f"- {achievement}")

            lines.append("")

    lines.append("EDUCATION")
    for edu in resume.get("education", []):
        if isinstance(edu, dict):
            lines.append(" | ".join([
                x for x in [
                    edu.get("degree", ""),
                    edu.get("university", ""),
                    edu.get("year", "")
                ] if x
            ]))

    lines.append("")
    lines.append("CERTIFICATIONS")
    for cert in resume.get("certifications", []):
        if isinstance(cert, dict):
            name = cert.get("name", "")
            desc = cert.get("description", "")
            lines.append(f"{name} - {desc}".strip(" -"))

    return "\n".join(lines)


def create_pdf_bytes(title, content):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", "B", 14)
    pdf.cell(0, 10, clean_text(title), ln=True)
    pdf.ln(5)

    pdf.set_font("Arial", size=11)

    if isinstance(content, dict):
        content = resume_to_plain_text(content)

    for line in str(content).splitlines():
        clean_line = clean_text(line)
        clean_line = clean_line.encode("latin-1", "replace").decode("latin-1")
        pdf.multi_cell(0, 6, clean_line)

    return pdf.output(dest="S").encode("latin-1", "ignore")


# ===========================
# DOCX GENERATOR
# ===========================
def create_docx_bytes(title, content):
    doc = Document()

    def add_centered(text, size=16, bold=True):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = bold
        run.font.size = Pt(size)

    def add_section_header(text):
        p = doc.add_paragraph()
        run = p.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(12)

        line = doc.add_paragraph()
        line_run = line.add_run("_" * 70)
        line_run.font.size = Pt(6)

    def add_text(text):
        if text:
            doc.add_paragraph(str(text))

    def add_bullet(text):
        if text:
            doc.add_paragraph(str(text), style="List Bullet")

    if isinstance(content, dict):
        name = content.get("name", "")
        title_value = content.get("title", "")
        contact = contact_line(content)

        if name:
            add_centered(name, size=18, bold=True)

        if title_value:
            add_centered(title_value, size=11, bold=False)

        if contact:
            add_centered(contact, size=10, bold=False)

        doc.add_paragraph("")

        add_section_header("Summary")
        add_text(content.get("summary", ""))
        doc.add_paragraph("")

        add_section_header("Technologies & Skills")
        skills = content.get("skills", [])

        if isinstance(skills, list):
            for skill_group in skills:
                if isinstance(skill_group, dict):
                    category = skill_group.get("category", "")
                    items = skill_group.get("items", [])
                    if category and items:
                        add_bullet(f"{category}: {', '.join(items)}")
                else:
                    add_bullet(skill_group)

        elif isinstance(skills, dict):
            for category, items in skills.items():
                label = category.replace("_", " ").title()
                if isinstance(items, list):
                    add_bullet(f"{label}: {', '.join(items)}")
                else:
                    add_bullet(f"{label}: {items}")

  
        doc.add_paragraph("")

        add_section_header("Projects")
        for project in content.get("projects", []):
            if isinstance(project, dict):
                name = project.get("name", "")
                description = project.get("description", "")
                achievements = clean_achievements(description, project.get("achievements", []))

                if name:
                    p = doc.add_paragraph()
                    run = p.add_run(name)
                    run.bold = True

                if description and not achievements:
                    add_text(description)

                for achievement in achievements:
                    add_bullet(achievement)

        doc.add_paragraph("")

        add_section_header("Education")
        for edu in content.get("education", []):
            if isinstance(edu, dict):
                text = " | ".join([
                    x for x in [
                        edu.get("degree", ""),
                        edu.get("university", ""),
                        edu.get("year", "")
                    ] if x
                ])
                add_bullet(text)

        doc.add_paragraph("")

        add_section_header("Certifications")
        for cert in content.get("certifications", []):
            if isinstance(cert, dict):
                name = cert.get("name", "")
                desc = cert.get("description", "")
                add_bullet(f"{name} - {desc}".strip(" -"))

    else:
        add_centered(title, size=16, bold=True)
        doc.add_paragraph("")
        add_text(str(content))

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def create_zip(files):
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w") as z:
        for name, data in files.items():
            z.writestr(name, data)
    buffer.seek(0)
    return buffer.getvalue()


# ===========================
# BUTTON
# ===========================

if st.button("Run AI Optimization 🚀"):

    if not resume_file or not cover_file or not job_requirements:
        st.error("Please fill all fields")
        st.stop()

    files = {
        "resume": (resume_file.name, resume_file, resume_file.type),
        "cover_letter": (cover_file.name, cover_file, cover_file.type)
    }

    data = {
        "job_requirements": job_requirements,
        "github_url": github_url
    }


    with st.spinner("🤖 AI agents are analyzing and optimizing your application..."):

        try:
            response = requests.post(
                f"{API_URL}/process-resume",
                files=files,
                data=data,
                timeout=300
            )
            
        except RequestException as e:
            st.error(f"Backend error: {e}")
            st.stop()

        if response.status_code != 200:
            st.error(response.text)
            st.stop()

        raw = response.json()
        st.session_state.result = raw.get("data", {})

    st.success("✅ Processing complete!")


# ===========================
# OUTPUT
# ===========================
result = st.session_state.result

if result:

    tabs = st.tabs([
        "📊 Analysis",
        "✏️ Editing",
        "📄 ATS",
        "🧠 Review",
        "⬇️ Downloads"
    ])

    with tabs[0]:
        analysis = result.get("analysis", {})

        st.markdown("## 📊 Analysis")

        section_header("Missing Keywords")
        st.markdown(bullet_list(analysis.get("missing_keywords", [])))

        section_header("Strengths")
        st.markdown(bullet_list(analysis.get("strengths", [])))

        section_header("Weaknesses")
        st.markdown(bullet_list(analysis.get("weaknesses", [])))

        section_header("Mismatches")
        st.markdown(bullet_list(analysis.get("mismatches", [])))

        section_header("Suggestions")
        st.markdown(bullet_list(analysis.get("improvement_suggestions", [])))

    with tabs[1]:
        editing = result.get("editing", {})
        resume = editing.get("resume", {})
        cover = editing.get("cover_letter", {})

        st.markdown("## 📄 Edited Resume")
        render_resume(resume)

        st.markdown("## ✉️ Cover Letter")
        st.write(cover.get("content", ""))

    with tabs[2]:
        ats = get_ats(result)
        ats_resume = ats.get("ats_resume", {})

        st.markdown("## 📄 ATS Resume")
        render_resume(ats_resume)

        st.markdown("## ✉️ ATS Cover Letter")
        st.write(ats.get("ats_cover_letter", ""))

    with tabs[3]:
        review = result.get("review", {})

        st.markdown("## 🧠 Review")

        st.markdown(f"""
- **ATS Score:** {review.get('ATS_score')}
- **Overall Score:** {review.get('overall_score')}
""")

        section_header("Strengths")
        st.markdown(bullet_list(review.get("strengths", [])))

        section_header("Weaknesses")
        st.markdown(bullet_list(review.get("weaknesses", [])))

        section_header("Suggestions")
        st.markdown(bullet_list(review.get("improvement_suggestions", [])))

    with tabs[4]:
        ats = get_ats(result)
        ats_resume = ats.get("ats_resume", {})
        ats_cover = ats.get("ats_cover_letter", "")

        resume_pdf = create_pdf_bytes("ATS Resume", ats_resume)
        cover_pdf = create_pdf_bytes("ATS Cover Letter", ats_cover)

        resume_docx = create_docx_bytes("ATS Resume", ats_resume)
        cover_docx = create_docx_bytes("ATS Cover Letter", ats_cover)

        zip_file = create_zip({
            "ats_resume.pdf": resume_pdf,
            "ats_cover_letter.pdf": cover_pdf,
            "ats_resume.docx": resume_docx,
            "ats_cover_letter.docx": cover_docx
        })

        st.download_button("📥 ATS Resume PDF", resume_pdf, "resume.pdf")
        st.download_button("📥 ATS Resume DOCX", resume_docx, "resume.docx")

        st.download_button("📥 ATS Cover Letter PDF", cover_pdf, "cover.pdf")
        st.download_button("📥 ATS Cover Letter DOCX", cover_docx, "cover.docx")

        st.download_button("📦 Download All Files (ZIP)", zip_file, "ats_package.zip")